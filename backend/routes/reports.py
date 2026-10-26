import csv
import io
from datetime import date
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from backend.database import get_db
from backend.dependencies import require_roles
from backend.models.user import User, UserRole
from backend.services.shift_fuel_report_service import ShiftFuelReportService


router = APIRouter(prefix="/api/reports", tags=["Reports"])

REPORT_COLUMNS = [
    ("Vessel Name", "vessel_name"),
    ("Shift Date", "shift_date"),
    ("Shift Time (Morning/Evening)", "shift_name"),
    ("Fuel Received Quantity (L)", "fuel_received_quantity"),
    ("Fuel Received From", "fuel_received_from"),
    ("Main Engine Start Time", "main_engine_start_time"),
    ("Main Engine Close Time", "main_engine_close_time"),
    ("Main Engine Total Running H/R", "main_engine_running_hours"),
    ("Main Engine Consumption (L)", "main_engine_consumption"),
    ("AUX Engine Start Time", "aux_engine_start_time"),
    ("AUX Engine Close Time", "aux_engine_close_time"),
    ("AUX Engine Total Running H/R", "aux_engine_running_hours"),
    ("AUX Engine Consumption (L)", "aux_engine_consumption"),
    ("DG Start Time", "dg_start_time"),
    ("DG Close Time", "dg_close_time"),
    ("DG Total Running H/R", "dg_running_hours"),
    ("DG Consumption (L)", "dg_consumption"),
    ("Total F.O Consumption of the Vessel (L)", "total_vessel_fo_consumption"),
    ("Transfer Out (L)", "transfer_out_quantity"),
    ("Transfer Out Vessel", "transfer_out_vessels"),
    ("Transfer In (L)", "transfer_in_quantity"),
    ("Transfer In Vessel", "transfer_in_vessels"),
    ("Opening Fuel (L)", "opening_fuel"),
    ("Closing Fuel (L)", "closing_fuel"),
    ("Advancement (m)", "advancement_m"),
    ("Dredging Hours", "dredging_hours"),
    ("Sounding / Fuel Level (L)", "sounding_fuel_level_litres"),
    ("Remarks", "remarks"),
]


def _get_report_rows(
    db: Session,
    from_date: date,
    to_date: date,
    vessel_id: int | None,
) -> list[dict]:
    if from_date > to_date:
        raise HTTPException(status_code=400, detail="From date cannot be later than to date.")
    return ShiftFuelReportService.get_rows(
        db,
        from_date=from_date,
        to_date=to_date,
        vessel_id=vessel_id,
    )


def _safe_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, date):
        value = str(value)
    else:
        value = str(value)
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


@router.get("/shift-fuel")
def get_shift_fuel_report(
    from_date: date = Query(...),
    to_date: date = Query(...),
    vessel_id: int | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.MANAGER, UserRole.ADMIN)),
):
    return _get_report_rows(db, from_date, to_date, vessel_id)


@router.get("/shift-fuel.csv")
def export_shift_fuel_report_csv(
    from_date: date = Query(...),
    to_date: date = Query(...),
    vessel_id: int | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.MANAGER, UserRole.ADMIN)),
):
    rows = _get_report_rows(db, from_date, to_date, vessel_id)
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([header for header, _ in REPORT_COLUMNS])
    for row in rows:
        writer.writerow([_safe_cell(row[field]) for _, field in REPORT_COLUMNS])
    filename = f"shift_fuel_report_{from_date}_{to_date}.csv"
    return Response(
        content="\ufeff" + output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/shift-fuel.xlsx")
def export_shift_fuel_report_excel(
    from_date: date = Query(...),
    to_date: date = Query(...),
    vessel_id: int | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.MANAGER, UserRole.ADMIN)),
):
    rows = _get_report_rows(db, from_date, to_date, vessel_id)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Shift Fuel Report"
    worksheet.append([header for header, _ in REPORT_COLUMNS])
    for row in rows:
        worksheet.append([_safe_cell(row[field]) for _, field in REPORT_COLUMNS])

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        letter = column_cells[0].column_letter
        max_length = max((len(str(cell.value or "")) for cell in column_cells), default=10)
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 38)
    worksheet.row_dimensions[1].height = 34

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"shift_fuel_report_{from_date}_{to_date}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
